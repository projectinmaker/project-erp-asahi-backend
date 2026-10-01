"""Helper generik untuk membuat workbook Excel (openpyxl) siap di-stream.

Update #5 — fitur Import & Export Excel.
Dipakai endpoint export master (barang/pelanggan/supplier), export tagihan
piutang/hutang, export stok, dan template import master.

Style minimal-profesional:
- Header: bold + fill abu terang + border tipis, freeze panes A2.
- Auto column width (perkiraan panjang isi, min 10 max 45).
- Kolom angka: number_format '#,##0' (integer) / '#,##0.00' (desimal);
  kolom tanggal: 'DD/MM/YYYY'.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="F2F2F2")
HEADER_FONT = Font(bold=True)
_THIN = Side(style="thin", color="D9D9D9")
THIN_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

INTEGER_FORMAT = "#,##0"
DECIMAL_FORMAT = "#,##0.00"
DATE_FORMAT = "DD/MM/YYYY"

MIN_WIDTH = 10
MAX_WIDTH = 45


def _cell_text(value) -> str:
    """Representasi teks sebuah nilai untuk estimasi lebar kolom."""
    if value is None:
        return ""
    return str(value)


def workbook_from_rows(
    headers: list[str],
    rows: list[list],
    sheet: str = "Data",
    number_columns: set[int] | None = None,
    decimal_columns: set[int] | None = None,
    date_columns: set[int] | None = None,
) -> Workbook:
    """Bangun Workbook dari header + baris data.

    number_columns  : indeks kolom (0-based) dengan format integer '#,##0'.
    decimal_columns : indeks kolom (0-based) dengan format desimal '#,##0.00'.
    date_columns    : indeks kolom (0-based) berisi date/datetime → 'DD/MM/YYYY'.
    """
    number_columns = number_columns or set()
    decimal_columns = decimal_columns or set()
    date_columns = date_columns or set()

    wb = Workbook()
    ws = wb.active
    ws.title = sheet

    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        ws.append(list(row))

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = THIN_BORDER

    # Number/date format per kolom (0-based index → 1-based kolom excel).
    for idx in sorted(number_columns | decimal_columns | date_columns):
        fmt = DECIMAL_FORMAT if idx in decimal_columns else INTEGER_FORMAT if idx in number_columns else DATE_FORMAT
        for row in ws.iter_rows(min_row=2, min_col=idx + 1, max_col=idx + 1):
            row[0].number_format = fmt

    # Auto column width: perkiraan panjang isi (header + data), min 10 max 45.
    for idx, header in enumerate(headers, start=1):
        longest = len(_cell_text(header))
        for row in rows:
            if idx - 1 < len(row):
                longest = max(longest, len(_cell_text(row[idx - 1])))
        ws.column_dimensions[get_column_letter(idx)].width = max(MIN_WIDTH, min(MAX_WIDTH, longest + 2))

    ws.freeze_panes = "A2"
    return wb


def add_instructions_sheet(wb: Workbook, rows: list[list], sheet: str = "Petunjuk") -> None:
    """Tambah sheet petunjuk (kolom, wajib, tipe data, keterangan + catatan umum).

    rows: list baris — baris pertama dianggap header sheet Petunjuk.
    """
    ws = wb.create_sheet(title=sheet)
    for row in rows:
        ws.append(list(row))
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if rows:
        n_cols = len(rows[0])
        for idx in range(1, n_cols + 1):
            longest = max(len(_cell_text(r[idx - 1])) for r in rows)
            # Kolom keterangan terakhir diberi lebar penuh agar catatan terbaca.
            width = MAX_WIDTH if idx == n_cols else max(MIN_WIDTH, min(MAX_WIDTH, longest + 2))
            ws.column_dimensions[get_column_letter(idx)].width = width


def workbook_to_stream(wb: Workbook) -> BytesIO:
    """Simpan workbook ke BytesIO siap dikirim via StreamingResponse."""
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def xlsx_response(wb: Workbook, filename: str) -> dict:
    """Metadata response standar untuk StreamingResponse file .xlsx."""
    return {
        "stream": workbook_to_stream(wb),
        "media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "headers": {"Content-Disposition": f'attachment; filename="{filename}"'},
    }
